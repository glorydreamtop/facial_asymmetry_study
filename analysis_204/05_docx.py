"""Build the paper-style Word document with the main tables and figures."""
from __future__ import annotations

import csv
import json
from pathlib import Path

import numpy as np
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

import common as C

OUT = C.OUT
DELIV = C.DELIV
FIG = DELIV / "figures"


def set_run_font(run, size=10, bold=False, name="Times New Roman", east="宋体"):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    set_run_font(r, size=14 if level == 1 else 12, bold=True, east="黑体")
    return p


def add_para(doc, text, size=9.5, italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    set_run_font(r, size=size, bold=False)
    r.italic = italic
    return p


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(text)
    set_run_font(r, size=8.5, bold=True)
    return p


def add_table(doc, headers, rows, widths=None, font_size=8):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        set_run_font(r, size=font_size, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            r = p.add_run(str(v))
            set_run_font(r, size=font_size)
    if widths:
        for i, w in enumerate(widths):
            for row in table.rows:
                row.cells[i].width = Inches(w)
    return table


def add_figure(doc, filename, caption, width=6.2):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(FIG / filename), width=Inches(width))
    add_caption(doc, caption)


def main():
    axis = json.load(open(OUT / "axis_final.json", encoding="utf-8"))
    final = json.load(open(OUT / "final_results.json", encoding="utf-8"))
    per = list(csv.DictReader(open(OUT / "final_per_case.csv", encoding="utf-8-sig")))
    test = [r for r in per if r["split"] == "test"]
    regions = ["eyebrow", "eye", "nose", "mouth", "contour"]
    labels = ["Eyebrow", "Eye", "Nose", "Mouth", "Facial contour"]

    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = Inches(0.8)
    sec.right_margin = Inches(0.8)
    sec.top_margin = Inches(0.8)
    sec.bottom_margin = Inches(0.8)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(9.5)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    # title
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("面部不对称自动评分：论文数据与图表")
    set_run_font(r, size=16, bold=True, east="黑体")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("基于 29 个面部关键点、专家共识对称轴与 Global FA 指数")
    set_run_font(r, size=10.5)

    d = final["descriptives"]
    add_heading(doc, "1. 数据与划分")
    add_para(doc, f"共标注 216 例，双侧关键点完整纳入 204 例；12 例因关键点缺失排除。训练 {d['split']['train']} 例、验证 {d['split']['val']} 例、测试 {d['split']['test']} 例（7:2:1，seed=11）。图像分辨率均为 1856×2784。IPD 为 {d['ipd']['mean']:.1f} ± {d['ipd']['sd']:.1f} px（中位数 {d['ipd']['median']:.1f}，范围 {d['ipd']['min']:.0f}–{d['ipd']['max']:.0f}）。专家 VAS 为 {d['vas']['mean']:.1f} ± {d['vas']['sd']:.1f}（范围 {d['vas']['min']:.0f}–{d['vas']['max']:.0f}）。")

    add_heading(doc, "2. 对称轴预测")
    add_para(doc, "输入为 29 个关键点构造的 167 维特征，输出为专家共识轴参数 (m, b)。模型在训练+验证集（n=184）上拟合，测试集（n=20）只评估一次。")
    rows = []
    for name in ["Ridge", "Linear", "SVR-RBF", "Random Forest"]:
        m = axis["models"][name]
        t = m["test"]
        rows.append([
            name,
            f"{m['cv']['mean']:.3f}",
            f"{t['angle']['mean']:.3f} ± {t['angle']['sd']:.3f}",
            f"{t['angle']['median']:.3f} ({t['angle']['q1']:.3f}–{t['angle']['q3']:.3f})",
            f"{t['position']['median']:.3f} ({t['position']['q1']:.3f}–{t['position']['q3']:.3f})",
            f"{t['mhd']['median']:.3f} ({t['mhd']['q1']:.3f}–{t['mhd']['q3']:.3f})",
        ])
    add_caption(doc, "Table 1. Comparison of symmetry-axis prediction errors among the four regression models.")
    add_table(doc, ["Model", "CV angle (°)", "Angle mean ± SD (°)", "Angle median (IQR)", "Position median (IQR)", "MHD median (IQR)"],
              rows, widths=[0.95, 0.8, 1.2, 1.35, 1.35, 1.25])
    add_para(doc, "Test set n=20. CV: 5-fold cross-validation on train+val (n=184). MHD: mean horizontal distance over 100 samples between y=500 and y=2280. Position is evaluated at the interpupillary-line height yc.", size=8, italic=True)

    t = axis["models"]["Ridge"]["test"]
    add_caption(doc, "Table 2. Agreement between the Ridge-predicted symmetry axis and the expert consensus axis.")
    rows = []
    for label, key, sk in [("Angular deviation (°)", "angle", "signed_angle"),
                           ("Axis-position deviation at yc (px)", "position", "signed_position")]:
        s = t[key]; sa = t[sk]
        rows.append([label, f"{s['mean']:.2f} ± {s['sd']:.2f}",
                     f"{s['median']:.2f} ({s['q1']:.2f}–{s['q3']:.2f})",
                     f"{s['ci_low']:.2f}–{s['ci_high']:.2f}",
                     f"{sa['bias']:.2f} ± {sa['sd']:.2f}",
                     f"{sa['loa_low']:.2f} to {sa['loa_high']:.2f}"])
    add_table(doc, ["Metric", "Mean ± SD", "Median (IQR)", "95% CI", "Bias ± SD", "95% LoA"],
              rows, widths=[1.5, 0.85, 1.15, 0.85, 0.85, 1.0])
    add_para(doc, "Test set n=20. CI: 95% confidence interval of the mean. LoA: limits of agreement (bias ± 1.96 SD).", size=8, italic=True)

    add_figure(doc, "fig1_axis_angle_boxplot.png", "Figure 1. Angular deviation of the four axis models (test).", 5.6)
    add_figure(doc, "fig2_axis_position_boxplot.png", "Figure 2. Axis-position deviation at yc (test).", 5.6)
    add_figure(doc, "fig3_ridge_bland_altman.png", "Figure 3. Bland–Altman plots for the Ridge axis (angle and position).", 6.4)

    add_heading(doc, "3. Global FA 公式")
    fit = final["formula"]
    add_para(doc, f"在训练+验证集（n={fit['n_fit']}）上拟合，R²={fit['r2']:.4f}，adjusted R²={fit['adj_r2']:.4f}，RMSE={fit['rmse']:.3f}。")
    pos = [max(0.0, x["beta_std"]) for x in fit["terms"]]
    wsum = sum(pos) if sum(pos) > 0 else 1.0
    rows = [["Intercept", f"{fit['intercept']:.3f} ({fit['intercept_ci'][0]:.3f}–{fit['intercept_ci'][1]:.3f})", "", "", ""]]
    for x, w in zip(fit["terms"], pos):
        rows.append([x["region"], f"{x['B']:.3f} ({x['B_ci_low']:.3f}–{x['B_ci_high']:.3f})",
                     f"{x['beta_std']:.3f} ({x['beta_std_ci_low']:.3f}–{x['beta_std_ci_high']:.3f})",
                     f"{x['p']:.2e}", f"{w/wsum:.3f}"])
    add_caption(doc, "Table 3. Multivariable linear regression model for the Global FA Index.")
    add_table(doc, ["Facial region", "Unstandardized B (95% CI)", "Standardized β (95% CI)", "P value", "Normalized weight"],
              rows, widths=[1.2, 1.7, 1.7, 0.8, 0.9])
    add_para(doc, "Global FA = min(100, max(0, β0 + Σ βr × RASr)). Normalized weight = positive standardized β / sum.", size=8, italic=True)
    ra = final["range_check_all"]
    add_caption(doc, "Range check before and after clipping.")
    rows = [["Raw z", f"{ra['raw_min']:.2f} / {ra['raw_max']:.2f}", ra["n_below_0"], ra["n_above_100"], f"{ra['out_of_range_pct']:.2f}"],
            ["Global FA (clipped)", f"{ra['fa_min']:.2f} / {ra['fa_max']:.2f}", 0, 0, "0.00"]]
    add_table(doc, ["Score", "Minimum / Maximum", "Below 0", "Above 100", "Out-of-range (%)"],
              rows, widths=[1.4, 1.8, 0.8, 0.9, 1.1])

    add_heading(doc, "4. 机器评分与专家 VAS")
    a = final["routes"]["A"]
    add_para(doc, f"路线 A：自动关键点 + 预测轴。Pearson r={a['pearson_r']:.3f}（95% CI {a['pearson_r_ci'][0]:.3f}–{a['pearson_r_ci'][1]:.3f}），R²={a['r2']:.3f}，ICC(A,1)={a['icc_a1']:.3f}（95% CI {a['icc_a1_ci'][0]:.3f}–{a['icc_a1_ci'][1]:.3f}），Bias={a['bias']:.2f}（95% LoA {a['loa_low']:.2f} 至 {a['loa_high']:.2f}），MAE={a['mae']:.2f}（95% CI {a['mae_ci'][0]:.2f}–{a['mae_ci'][1]:.2f}），RMSE={a['rmse']:.2f}（95% CI {a['rmse_ci'][0]:.2f}–{a['rmse_ci'][1]:.2f}）。")
    add_caption(doc, "Table 4. Agreement between the Global FA Index and expert VAS scores (route A).")
    rows = [["Pearson correlation", f"{a['pearson_r']:.3f}", f"{a['pearson_r_ci'][0]:.3f}–{a['pearson_r_ci'][1]:.3f}", "<0.001"],
            ["Coefficient of determination", f"{a['r2']:.3f}", "—", "—"],
            ["ICC(A,1)", f"{a['icc_a1']:.3f}", f"{a['icc_a1_ci'][0]:.3f}–{a['icc_a1_ci'][1]:.3f}", "<0.001"],
            ["Bland–Altman bias", f"{a['bias']:.3f}", f"LoA {a['loa_low']:.3f} to {a['loa_high']:.3f}", "—"],
            ["MAE", f"{a['mae']:.3f}", f"{a['mae_ci'][0]:.3f}–{a['mae_ci'][1]:.3f}", "—"],
            ["RMSE", f"{a['rmse']:.3f}", f"{a['rmse_ci'][0]:.3f}–{a['rmse_ci'][1]:.3f}", "—"]]
    add_table(doc, ["Analysis method", "Estimate", "95% CI / LoA", "P value"], rows, widths=[1.6, 1.0, 1.6, 0.7])
    add_figure(doc, "fig4_fa_vs_vas.png", "Figure 4. Global FA vs expert VAS (test).", 4.3)
    add_figure(doc, "fig5_fa_residual.png", "Figure 5. Residual plot (Global FA − VAS).", 4.6)
    add_figure(doc, "fig6_fa_bland_altman.png", "Figure 6. Bland–Altman plot for Global FA vs expert VAS.", 4.8)
    add_figure(doc, "fig7_ras_distribution.png", "Figure 7. Regional asymmetry scores (test).", 6.2)

    add_heading(doc, "5. 补充实验一：三条路线")
    ax = final["axis_errors_test"]
    add_caption(doc, "Panel A. Downstream axis errors (test, median (IQR)).")
    rows = []
    for key, label in [("A", "A axis vs consensus"), ("B", "B axis vs consensus"), ("A_vs_B", "A axis vs B axis")]:
        s = ax[key]
        rows.append([label, f"{s['angle']['median']:.3f} ({s['angle']['q1']:.3f}–{s['angle']['q3']:.3f})",
                     f"{s['position']['median']:.3f} ({s['position']['q1']:.3f}–{s['position']['q3']:.3f})",
                     f"{s['mhd']['median']:.3f} ({s['mhd']['q1']:.3f}–{s['mhd']['q3']:.3f})"])
    add_table(doc, ["Comparison", "Angle (°)", "Position (px)", "MHD (px)"], rows, widths=[1.7, 1.5, 1.5, 1.4])

    add_caption(doc, "Panel B. Regional and global score differences (test, mean ± SD).")
    rows = []
    for key, label in zip(regions, labels):
        av = np.array([float(r[f"RAS_A_{key}"]) for r in test])
        bv = np.array([float(r[f"RAS_B_{key}"]) for r in test])
        cv = np.array([float(r[f"RAS_C_{key}"]) for r in test])
        rows.append([label, f"{av.mean():.4f} ± {av.std(ddof=1):.4f}", f"{bv.mean():.4f} ± {bv.std(ddof=1):.4f}",
                     f"{cv.mean():.4f} ± {cv.std(ddof=1):.4f}", f"{(av-bv).mean():.4f} ± {(av-bv).std(ddof=1):.4f}",
                     f"{(bv-cv).mean():.4f} ± {(bv-cv).std(ddof=1):.4f}"])
    fa = np.array([float(r["FA_A"]) for r in test]); fb = np.array([float(r["FA_B"]) for r in test]); fc = np.array([float(r["FA_C"]) for r in test])
    rows.append(["Global FA", f"{fa.mean():.2f} ± {fa.std(ddof=1):.2f}", f"{fb.mean():.2f} ± {fb.std(ddof=1):.2f}",
                 f"{fc.mean():.2f} ± {fc.std(ddof=1):.2f}", f"{(fa-fb).mean():.2f} ± {(fa-fb).std(ddof=1):.2f}",
                 f"{(fb-fc).mean():.2f} ± {(fb-fc).std(ddof=1):.2f}"])
    add_table(doc, ["Outcome", "A", "B", "C", "A − B", "B − C"], rows, widths=[1.0, 1.0, 1.0, 1.0, 1.0, 1.0], font_size=7.5)

    add_caption(doc, "Panel C. Global FA accuracy against expert VAS (test).")
    rows = []
    for key in ("A", "B", "C"):
        r = final["routes"][key]
        rows.append([key, f"{r['mae']:.2f} ({r['mae_ci'][0]:.2f}–{r['mae_ci'][1]:.2f})",
                     f"{r['rmse']:.2f} ({r['rmse_ci'][0]:.2f}–{r['rmse_ci'][1]:.2f})",
                     f"{r['r2']:.3f}", f"{r['pearson_r']:.3f}",
                     f"{r['icc_a1']:.3f} ({r['icc_a1_ci'][0]:.3f}–{r['icc_a1_ci'][1]:.3f})",
                     f"{r['bias']:.2f} ({r['loa_low']:.2f}–{r['loa_high']:.2f})"])
    add_table(doc, ["Route", "MAE (95% CI)", "RMSE (95% CI)", "R²", "r", "ICC(A,1)", "Bias (95% LoA)"],
              rows, widths=[0.45, 1.0, 1.0, 0.5, 0.45, 1.1, 1.0], font_size=7.5)
    add_para(doc, "ΔMAE(A−B) = −0.66 (95% CI −1.54 to 0.33), paired bootstrap 2000 resamples.", size=8, italic=True)

    add_heading(doc, "6. 补充实验二：AI 轴与单条专家轴")
    add_caption(doc, "Table S2. AI axis versus each individual expert axis, and inter-expert agreement (test, median (IQR)).")
    rows = []
    for label, s in final["supplement2"].items():
        rows.append([label, f"{s['angle']['median']:.3f} ({s['angle']['q1']:.3f}–{s['angle']['q3']:.3f})",
                     f"{s['position']['median']:.3f} ({s['position']['q1']:.3f}–{s['position']['q3']:.3f})",
                     f"{s['mhd']['median']:.3f} ({s['mhd']['q1']:.3f}–{s['mhd']['q3']:.3f})"])
    add_table(doc, ["Comparison", "Angle (°)", "Position (px)", "MHD (px)"], rows, widths=[1.7, 1.5, 1.5, 1.4])
    add_figure(doc, "fig8_supp2_axis.png", "Figure 8. AI axis vs individual expert axes, and inter-expert agreement (test).", 6.2)

    add_heading(doc, "7. 说明")
    add_para(doc, "12 例因双侧关键点不完整排除，204 例用于全部分析。评分时轴线以双瞳中点为锚点做横向平移；轴误差表仍使用原始轴线。Global FA 截断至 0–100，2/204（1.0%）例在截断前高于 100，已在越界检查中报告。167 维特征的定义、归一化、反变换、热图与训练参数见《167维特征说明》。")

    cp = doc.core_properties
    cp.author = ""
    cp.last_modified_by = ""
    cp.comments = ""
    cp.category = ""
    cp.keywords = ""
    cp.title = "面部不对称自动评分：论文数据与图表"
    out = DELIV / "论文数据与图表.docx"
    doc.save(out)
    print("saved", out)


if __name__ == "__main__":
    main()
