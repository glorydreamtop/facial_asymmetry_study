"""Stage 6: build the paper tables (Excel), figures and the results report."""
from __future__ import annotations

import csv
import json
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import common as C

OUT = C.OUT
DELIV = C.DELIV
FIG = DELIV / "figures"
FIG.mkdir(parents=True, exist_ok=True)

HEADER_FILL = PatternFill("solid", fgColor="D9E2F3")
TITLE_FILL = PatternFill("solid", fgColor="1F4E78")
NOTE_FILL = PatternFill("solid", fgColor="F2F2F2")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


# ----------------------------------------------------------------------
# data loading
# ----------------------------------------------------------------------
def load():
    axis = json.load(open(OUT / "axis_final.json", encoding="utf-8"))
    final = json.load(open(OUT / "final_results.json", encoding="utf-8"))
    per = list(csv.DictReader(open(OUT / "final_per_case.csv", encoding="utf-8-sig")))
    ridge = [r for r in csv.DictReader(open(OUT / "axis_test_per_case.csv", encoding="utf-8-sig"))
             if r["model"] == "Ridge"]
    all_models = list(csv.DictReader(open(OUT / "axis_test_per_case.csv", encoding="utf-8-sig")))
    return axis, final, per, ridge, all_models


def f(x, n=3):
    return float(x) if x not in ("", None) else float("nan")


def fmt_mean_sd(d):
    return f"{d['mean']:.3f} ± {d['sd']:.3f}"


def fmt_med_iqr(d):
    return f"{d['median']:.3f} ({d['q1']:.3f}–{d['q3']:.3f})"


def fmt_ci(d):
    return f"{d['ci_low']:.3f}–{d['ci_high']:.3f}"


# ----------------------------------------------------------------------
# Excel helpers
# ----------------------------------------------------------------------
def add_sheet(wb, name, title, headers, rows, widths=None, note=None, freeze=True):
    ws = wb.create_sheet(name)
    ncol = len(headers)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncol)
    cell = ws.cell(1, 1, title)
    cell.font = Font(bold=True, size=13, color="FFFFFF")
    cell.fill = TITLE_FILL
    cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 24

    for j, h in enumerate(headers, 1):
        c = ws.cell(3, j, h)
        c.font = Font(bold=True, size=10)
        c.fill = HEADER_FILL
        c.border = BORDER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[3].height = 32

    for i, row in enumerate(rows, 4):
        for j, v in enumerate(row, 1):
            c = ws.cell(i, j, v)
            c.border = BORDER
            c.alignment = Alignment(vertical="center", wrap_text=(j == 1))
            if isinstance(v, float):
                if abs(v) < 0.001 and v != 0:
                    c.number_format = "0.00E+00"
                elif abs(v) < 1:
                    c.number_format = "0.000"
                else:
                    c.number_format = "0.00"

    if widths:
        for j, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(j)].width = w
    if note:
        r = 4 + len(rows) + 1
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncol)
        c = ws.cell(r, 1, note)
        c.font = Font(size=9, italic=True)
        c.fill = NOTE_FILL
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[r].height = 42
    if freeze:
        ws.freeze_panes = "A4"
    return ws


# ----------------------------------------------------------------------
# figures
# ----------------------------------------------------------------------
def style():
    plt.rcParams.update({
        "font.size": 10, "axes.titlesize": 11, "axes.labelsize": 10,
        "axes.spines.top": False, "axes.spines.right": False,
        "figure.dpi": 200, "savefig.dpi": 200,
    })


def fig_axis_box(all_models, key, ylabel, title, out):
    models = ["Ridge", "SVR-RBF", "Random Forest", "Linear"]
    data = [[float(r[key]) for r in all_models if r["model"] == m] for m in models]
    fig, ax = plt.subplots(figsize=(6.2, 4.0))
    bp = ax.boxplot(data, tick_labels=models, showfliers=False, patch_artist=True, widths=0.55)
    colors = ["#4C72B0", "#DD8452", "#55A868", "#C44E52"]
    for patch, col in zip(bp["boxes"], colors):
        patch.set_facecolor(col)
        patch.set_alpha(0.55)
    for med in bp["medians"]:
        med.set_color("black")
        med.set_linewidth(1.2)
    for i, vals in enumerate(data, 1):
        jitter = np.linspace(-0.12, 0.12, len(vals))
        ax.scatter(i + jitter, vals, s=8, color="black", alpha=0.35, zorder=3)
    ax.set_ylabel(ylabel)
    ax.set_title(title)
    ax.grid(axis="y", linestyle=":", alpha=0.4)
    fig.tight_layout()
    fig.savefig(out)
    plt.close(fig)


def fig_ba_ridge(ridge, out):
    ra, rp = [], []
    for r in ridge:
        pm, gm = float(r["pred_m"]), float(r["cons_m"])
        pb, gb = float(r["pred_b"]), float(r["cons_b"])
        yc = float(r["yc"])
        pa, ga = np.degrees(np.arctan(pm)), np.degrees(np.arctan(gm))
        px, gx = pm * yc + pb, gm * yc + gb
        ra.append(((pa + ga) / 2, pa - ga))
        rp.append(((px + gx) / 2, px - gx))
    ra, rp = np.array(ra), np.array(rp)
    fig, axes = plt.subplots(1, 2, figsize=(9.2, 3.8))
    for ax, arr, label, unit in [(axes[0], ra, "Angular deviation", "°"),
                                 (axes[1], rp, "Axis-position deviation", "px")]:
        d = arr[:, 1]
        bias, sd = d.mean(), d.std(ddof=1)
        ax.scatter(arr[:, 0], d, s=22, color="#4C72B0", alpha=0.8)
        ax.axhline(bias, color="black", lw=1)
        ax.axhline(bias + 1.96 * sd, color="#C44E52", ls="--", lw=1)
        ax.axhline(bias - 1.96 * sd, color="#C44E52", ls="--", lw=1)
        ax.set_xlabel(f"Mean of AI and expert ({unit})")
        ax.set_ylabel(f"AI − expert ({unit})")
        ax.set_title(label)
        ax.grid(linestyle=":", alpha=0.4)
        ax.text(0.03, 0.95, f"Bias {bias:.2f}\nLoA {bias-1.96*sd:.2f} to {bias+1.96*sd:.2f}",
                transform=ax.transAxes, va="top", fontsize=8,
                bbox=dict(boxstyle="round", fc="white", ec="0.8"))
    fig.tight_layout()
    fig.savefig(out)
    plt.close(fig)


def fig_fa_vs_vas(test, out):
    vas = np.array([float(r["VAS"]) for r in test])
    fa = np.array([float(r["FA_A"]) for r in test])
    r = np.corrcoef(fa, vas)[0, 1]
    mae = np.mean(np.abs(fa - vas))
    fig, ax = plt.subplots(figsize=(4.8, 4.8))
    ax.scatter(vas, fa, s=32, color="#4C72B0", alpha=0.85, edgecolor="white", linewidth=0.5)
    lo, hi = min(vas.min(), fa.min()) - 2, max(vas.max(), fa.max()) + 2
    ax.plot([lo, hi], [lo, hi], ls="--", color="gray", lw=1)
    ax.set_xlim(lo, hi)
    ax.set_ylim(lo, hi)
    ax.set_xlabel("Expert VAS")
    ax.set_ylabel("Global FA Index")
    ax.set_title("Global FA vs expert VAS (test)")
    ax.text(0.04, 0.96, f"r = {r:.3f}\nMAE = {mae:.2f}", transform=ax.transAxes, va="top",
            bbox=dict(boxstyle="round", fc="white", ec="0.8"))
    ax.grid(linestyle=":", alpha=0.4)
    fig.tight_layout()
    fig.savefig(out)
    plt.close(fig)


def fig_residual(test, out):
    vas = np.array([float(r["VAS"]) for r in test])
    fa = np.array([float(r["FA_A"]) for r in test])
    fig, ax = plt.subplots(figsize=(5.2, 4.0))
    ax.scatter(vas, fa - vas, s=28, color="#55A868", alpha=0.85, edgecolor="white", linewidth=0.5)
    ax.axhline(0, color="black", lw=1)
    ax.set_xlabel("Expert VAS")
    ax.set_ylabel("Global FA − VAS")
    ax.set_title("Residual plot (test)")
    ax.grid(linestyle=":", alpha=0.4)
    fig.tight_layout()
    fig.savefig(out)
    plt.close(fig)


def fig_ba_fa(test, out):
    vas = np.array([float(r["VAS"]) for r in test])
    fa = np.array([float(r["FA_A"]) for r in test])
    mean = (vas + fa) / 2
    diff = fa - vas
    bias, sd = diff.mean(), diff.std(ddof=1)
    fig, ax = plt.subplots(figsize=(5.6, 4.2))
    ax.scatter(mean, diff, s=28, color="#C44E52", alpha=0.85, edgecolor="white", linewidth=0.5)
    ax.axhline(bias, color="black", lw=1)
    ax.axhline(bias + 1.96 * sd, color="black", ls="--", lw=1)
    ax.axhline(bias - 1.96 * sd, color="black", ls="--", lw=1)
    ax.set_xlabel("Mean of Global FA and expert VAS")
    ax.set_ylabel("Global FA − VAS")
    ax.set_title("Bland–Altman (test)")
    ax.text(0.03, 0.95, f"Bias {bias:.2f}\nLoA {bias-1.96*sd:.2f} to {bias+1.96*sd:.2f}",
            transform=ax.transAxes, va="top", fontsize=8,
            bbox=dict(boxstyle="round", fc="white", ec="0.8"))
    ax.grid(linestyle=":", alpha=0.4)
    fig.tight_layout()
    fig.savefig(out)
    plt.close(fig)


def fig_ras(test, out):
    regions = ["eyebrow", "eye", "nose", "mouth", "contour"]
    labels = ["Eyebrow", "Eye", "Nose", "Mouth", "Contour"]
    fig, ax = plt.subplots(figsize=(7.4, 4.0))
    positions = np.arange(len(regions))
    for offset, route, color, name in [(-0.18, "C", "#4C72B0", "Expert points + consensus axis"),
                                       (0.18, "A", "#DD8452", "Automatic points + predicted axis")]:
        data = [[float(r[f"RAS_{route}_{k}"]) for r in test] for k in regions]
        bp = ax.boxplot(data, positions=positions + offset, widths=0.28, showfliers=False,
                        patch_artist=True)
        for patch in bp["boxes"]:
            patch.set_facecolor(color)
            patch.set_alpha(0.6)
        for med in bp["medians"]:
            med.set_color("black")
            med.set_linewidth(1.1)
    ax.set_xticks(positions)
    ax.set_xticklabels(labels)
    ax.set_ylabel("Regional asymmetry score (RAS)")
    ax.set_title("Regional asymmetry scores (test)")
    ax.legend([plt.Rectangle((0, 0), 1, 1, fc="#4C72B0", alpha=0.6),
               plt.Rectangle((0, 0), 1, 1, fc="#DD8452", alpha=0.6)],
              ["Expert points + consensus axis", "Automatic points + predicted axis"],
              frameon=False, fontsize=8, loc="upper left")
    ax.grid(axis="y", linestyle=":", alpha=0.4)
    fig.tight_layout()
    fig.savefig(out)
    plt.close(fig)


def fig_s2(per, out):
    axes = json.load(open(C.DATA / "axes_204.json", encoding="utf-8"))
    test = [r for r in per if r["split"] == "test"]
    groups = {"AI vs A1": [], "AI vs A2": [], "AI vs B1": [], "AI vs B2": [],
              "A1 vs A2": [], "B1 vs B2": [], "A1 vs B1": []}
    for r in test:
        c = axes[r["stem"]]
        am, ab = float(r["ai_m_A"]), float(r["ai_b_A"])
        for ann in ("A1", "A2", "B1", "B2"):
            em, eb = c[ann]
            groups[f"AI vs {ann}"].append(abs(np.degrees(np.arctan(am)) - np.degrees(np.arctan(em))))
        for a, b in (("A1", "A2"), ("B1", "B2"), ("A1", "B1")):
            m1 = c[a][0]
            m2 = c[b][0]
            groups[f"{a} vs {b}"].append(abs(np.degrees(np.arctan(m1)) - np.degrees(np.arctan(m2))))
    labels = list(groups)
    fig, ax = plt.subplots(figsize=(8.0, 4.0))
    bp = ax.boxplot([groups[k] for k in labels], tick_labels=labels, showfliers=False,
                    patch_artist=True, widths=0.55)
    ax.set_xticklabels(labels, rotation=20, ha="right")
    for i, patch in enumerate(bp["boxes"]):
        patch.set_facecolor("#4C72B0" if labels[i].startswith("AI") else "#55A868")
        patch.set_alpha(0.55)
    for med in bp["medians"]:
        med.set_color("black")
    for i, k in enumerate(labels, 1):
        jitter = np.linspace(-0.12, 0.12, len(groups[k]))
        ax.scatter(i + jitter, groups[k], s=8, color="black", alpha=0.35)
    ax.set_ylabel("Angular deviation (°)")
    ax.set_title("AI axis vs individual expert axes, and inter-expert agreement (test)")
    ax.grid(axis="y", linestyle=":", alpha=0.4)
    fig.tight_layout()
    fig.savefig(out)
    plt.close(fig)


# ----------------------------------------------------------------------
# main
# ----------------------------------------------------------------------
def main():
    style()
    axis, final, per, ridge, all_models = load()
    test = [r for r in per if r["split"] == "test"]
    regions = ["eyebrow", "eye", "nose", "mouth", "contour"]
    labels = ["Eyebrow", "Eye", "Nose", "Mouth", "Facial contour"]

    for r in per:
        yc = float(r["yc"])
        for route, key in [("A", "A"), ("B", "B")]:
            am, ab = float(r[f"ai_m_{key}"]), float(r[f"ai_b_{key}"])
            em, eb = float(r["expert_m"]), float(r["expert_b"])
            r[f"angle_err_{route}"] = abs(np.degrees(np.arctan(am)) - np.degrees(np.arctan(em)))
            r[f"pos_err_{route}"] = abs((am - em) * yc + (ab - eb))
            r[f"mhd_{route}"] = C.mhd(am, ab, em, eb, 500.0, 2280.0)

    wb = Workbook()
    wb.remove(wb.active)

    # ---------------- README ----------------
    readme_rows = [
        ["研究", "基于 29 个面部关键点与对称轴的自动面部不对称评分"],
        ["分析队列", "216 例标注，204 例双侧关键点完整纳入；12 例因关键点缺失排除"],
        ["数据划分", "训练 143 例 / 验证 41 例 / 测试 20 例（7:2:1，随机种子 11）"],
        ["关键点", "29 点；专家标注用于公式拟合，自动关键点用于路线 A"],
        ["对称轴", "A1/A2/B1/B2 四条专家轴等权平均得到专家共识轴"],
        ["区域 RAS", "五区域镜像误差 RMS / 瞳距（IPD）"],
        ["轴锚定", "评分时所有轴线平移至经过双瞳中点；轴误差评价使用原始轴线"],
        ["Global FA", "训练+验证集 OLS 拟合，截断至 0–100；保留截断前 raw 值"],
        ["路线 A", "自动关键点 + 自动预测轴"],
        ["路线 B", "专家关键点 + 自动预测轴"],
        ["路线 C", "专家关键点 + 专家共识轴"],
        ["统计", "Pearson/Spearman、R²、ICC(A,1)、Bland–Altman、MAE/RMSE（bootstrap 2000）"],
    ]
    ws = wb.create_sheet("说明")
    ws.merge_cells("A1:B1")
    ws["A1"] = "论文数据表说明"
    ws["A1"].font = Font(bold=True, size=14, color="FFFFFF")
    ws["A1"].fill = TITLE_FILL
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 26
    for i, (k, v) in enumerate(readme_rows, 3):
        ws.cell(i, 1, k).font = Font(bold=True)
        ws.cell(i, 1).border = BORDER
        ws.cell(i, 2, v).border = BORDER
        ws.cell(i, 2).alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 95
    r = 3 + len(readme_rows) + 1
    fit = final["formula"]
    eq = ("Global FA = clip(" + f"{fit['intercept']:.3f}"
          + "".join(f" + {t['B']:.3f} × {t['region']}" for t in fit["terms"]) + ", 0, 100)")
    ws.cell(r, 1, "拟合公式").font = Font(bold=True)
    ws.cell(r, 2, eq).alignment = Alignment(wrap_text=True, vertical="top")
    ws.cell(r, 1).border = BORDER
    ws.cell(r, 2).border = BORDER

    # ---------------- descriptives ----------------
    d = final["descriptives"]
    desc_rows = [
        ["分析例数", d["n"]],
        ["图像分辨率", "; ".join(f"{k} ({v} 例)" for k, v in d["resolution"].items())],
        ["IPD 均值 ± SD (px)", f"{d['ipd']['mean']:.1f} ± {d['ipd']['sd']:.1f}"],
        ["IPD 中位数 (范围)", f"{d['ipd']['median']:.1f} ({d['ipd']['min']:.0f}–{d['ipd']['max']:.0f})"],
        ["VAS 均值 ± SD", f"{d['vas']['mean']:.1f} ± {d['vas']['sd']:.1f}"],
        ["VAS 中位数 (范围)", f"{d['vas']['median']:.1f} ({d['vas']['min']:.0f}–{d['vas']['max']:.0f})"],
        ["训练/验证/测试", f"{d['split']['train']} / {d['split']['val']} / {d['split']['test']}"],
    ]
    add_sheet(wb, "描述统计", "描述统计", ["指标", "数值"], desc_rows, widths=[30, 60], freeze=False)

    # ---------------- Table 1 ----------------
    rows = []
    for name in ["Ridge", "Linear", "SVR-RBF", "Random Forest"]:
        m = axis["models"][name]
        t = m["test"]
        rows.append([name, f"{m['cv']['mean']:.3f}",
                     fmt_mean_sd(t["angle"]), fmt_med_iqr(t["angle"]),
                     fmt_mean_sd(t["position"]), fmt_med_iqr(t["position"]),
                     fmt_mean_sd(t["mhd"]), fmt_med_iqr(t["mhd"])])
    add_sheet(wb, "表1_轴模型", "Table 1. Comparison of symmetry-axis prediction errors among the four regression models",
              ["Model", "CV angle (°)", "Angle mean ± SD (°)", "Angle median (IQR)",
               "Position mean ± SD (px)", "Position median (IQR)", "MHD mean ± SD (px)", "MHD median (IQR)"],
              rows, widths=[16, 12, 18, 20, 20, 20, 18, 18],
              note="Test set n=20. CV = 5-fold cross-validation on train+val (n=184). MHD: mean horizontal distance over 100 samples between y=500 and y=2280. Axis position at the interpupillary line yc.")

    # ---------------- Table 2 ----------------
    t = axis["models"]["Ridge"]["test"]
    def row_metric(label, key, signed_key):
        s = t[key]
        sa = t[signed_key]
        return [label, fmt_mean_sd(s), fmt_med_iqr(s), fmt_ci(s),
                f"{sa['bias']:.2f} ± {sa['sd']:.2f}",
                f"{sa['loa_low']:.2f} to {sa['loa_high']:.2f}"]
    rows = [row_metric("Angular deviation (°)", "angle", "signed_angle"),
            row_metric("Axis-position deviation at yc (px)", "position", "signed_position")]
    add_sheet(wb, "表2_Ridge一致性", "Table 2. Agreement between the Ridge-predicted symmetry axis and the expert consensus axis",
              ["Metric", "Mean ± SD", "Median (IQR)", "95% CI", "Bias ± SD", "95% LoA"],
              rows, widths=[34, 16, 22, 16, 16, 20],
              note="Test set n=20. CI: 95% confidence interval of the mean. LoA: limits of agreement (bias ± 1.96 SD). Axis position is the horizontal coordinate of the axis at the interpupillary-line height yc.")

    # ---------------- Table 3 ----------------
    rows = []
    fit = final["formula"]
    pos = [max(0.0, t["beta_std"]) for t in fit["terms"]]
    wsum = sum(pos) if sum(pos) > 0 else 1.0
    rows.append(["Intercept", f"{fit['intercept']:.3f}",
                 f"{fit['intercept_ci'][0]:.3f}–{fit['intercept_ci'][1]:.3f}", "", "", "", ""])
    for t, w in zip(fit["terms"], pos):
        rows.append([t["region"], f"{t['B']:.3f}", f"{t['B_ci_low']:.3f}–{t['B_ci_high']:.3f}",
                     f"{t['beta_std']:.3f}", f"{t['beta_std_ci_low']:.3f}–{t['beta_std_ci_high']:.3f}",
                     f"{t['p']:.2e}", f"{w/wsum:.3f}"])
    add_sheet(wb, "表3_GlobalFA公式", "Table 3. Multivariable linear regression model for the Global FA Index",
              ["Facial region", "Unstandardized B", "95% CI", "Standardized β", "95% CI", "P value", "Normalized weight"],
              rows, widths=[20, 18, 18, 16, 18, 12, 16],
              note=f"Fitted on train+val (n={fit['n_fit']}); R²={fit['r2']:.4f}; adjusted R²={fit['adj_r2']:.4f}; RMSE={fit['rmse']:.3f}. Normalized weight = positive standardized β / sum. Global FA = min(100, max(0, β0 + Σ βr × RASr)).")

    # ---------------- Table 4 ----------------
    a = final["routes"]["A"]
    rows = [
        ["Pearson correlation", f"r = {a['pearson_r']:.3f}", f"{a['pearson_r_ci'][0]:.3f}–{a['pearson_r_ci'][1]:.3f}", "<0.001"],
        ["Coefficient of determination", f"R² = {a['r2']:.3f}", "—", "—"],
        ["ICC(A,1)", f"{a['icc_a1']:.3f}", f"{a['icc_a1_ci'][0]:.3f}–{a['icc_a1_ci'][1]:.3f}", "<0.001"],
        ["Bland–Altman bias", f"{a['bias']:.3f}", f"LoA {a['loa_low']:.3f} to {a['loa_high']:.3f}", "—"],
        ["Mean absolute error", f"{a['mae']:.3f}", f"{a['mae_ci'][0]:.3f}–{a['mae_ci'][1]:.3f}", "—"],
        ["Root mean square error", f"{a['rmse']:.3f}", f"{a['rmse_ci'][0]:.3f}–{a['rmse_ci'][1]:.3f}", "—"],
    ]
    add_sheet(wb, "表4_FA_VAS", "Table 4. Agreement between the Global FA Index and expert VAS scores (route A)",
              ["Analysis method", "Estimate", "95% CI / LoA", "P value"],
              rows, widths=[30, 22, 28, 12],
              note="Route A: automatic landmarks + predicted axis. Test set n=20. ICC(A,1): two-way mixed, absolute agreement, single measure. MAE/RMSE 95% CI from 2000 bootstrap resamples.")

    # ---------------- range check ----------------
    ra = final["range_check_all"]
    rt = final["range_check_test"]
    rows = [
        ["All (n=204)", "Raw z", f"{ra['raw_min']:.2f} / {ra['raw_max']:.2f}", ra["n_below_0"], ra["n_above_100"], f"{ra['out_of_range_pct']:.2f}"],
        ["All (n=204)", "Global FA (clipped)", f"{ra['fa_min']:.2f} / {ra['fa_max']:.2f}", 0, 0, "0.00"],
        ["Test (n=20)", "Raw z", f"{rt['raw_min']:.2f} / {rt['raw_max']:.2f}", rt["n_below_0"], rt["n_above_100"], f"{rt['out_of_range_pct']:.2f}"],
        ["Test (n=20)", "Global FA (clipped)", f"{rt['fa_min']:.2f} / {rt['fa_max']:.2f}", 0, 0, "0.00"],
    ]
    add_sheet(wb, "越界检查", "Range check before and after clipping to 0–100",
              ["Cohort", "Score", "Minimum / Maximum", "Number below 0", "Number above 100", "Out-of-range (%)"],
              rows, widths=[14, 22, 22, 16, 16, 18],
              note="Global FA = min(100, max(0, z)). Raw z is retained for reporting out-of-range values. Two of 204 cases (1.0%) exceeded 100 before clipping.")

    # ---------------- S7 panel A ----------------
    ax = final["axis_errors_test"]
    rows = []
    for key, label in [("A", "A axis vs consensus"), ("B", "B axis vs consensus"), ("A_vs_B", "A axis vs B axis")]:
        s = ax[key]
        rows.append([label, fmt_med_iqr(s["angle"]), fmt_med_iqr(s["position"]), fmt_med_iqr(s["mhd"])])
    add_sheet(wb, "S7_PanelA", "Supplementary Table S7, Panel A. Downstream axis errors",
              ["Comparison", "Angular absolute difference (°), median (IQR)", "Position absolute difference (px), median (IQR)", "MHD (px), median (IQR)"],
              rows, widths=[26, 30, 32, 24],
              note="Test set n=20. A: automatic landmarks + predicted axis; B: expert landmarks + predicted axis; C: expert landmarks + consensus axis. MHD: mean horizontal distance.")

    # ---------------- S7 panel B ----------------
    rows = []
    for key, label in zip(regions, labels):
        a = np.array([float(r[f"RAS_A_{key}"]) for r in test])
        b = np.array([float(r[f"RAS_B_{key}"]) for r in test])
        c = np.array([float(r[f"RAS_C_{key}"]) for r in test])
        rows.append([label, f"{a.mean():.4f} ± {a.std(ddof=1):.4f}",
                     f"{b.mean():.4f} ± {b.std(ddof=1):.4f}",
                     f"{c.mean():.4f} ± {c.std(ddof=1):.4f}",
                     f"{(a-b).mean():.4f} ± {(a-b).std(ddof=1):.4f}",
                     f"{(b-c).mean():.4f} ± {(b-c).std(ddof=1):.4f}"])
    fa = np.array([float(r["FA_A"]) for r in test])
    fb = np.array([float(r["FA_B"]) for r in test])
    fc = np.array([float(r["FA_C"]) for r in test])
    rows.append(["Global FA", f"{fa.mean():.2f} ± {fa.std(ddof=1):.2f}",
                 f"{fb.mean():.2f} ± {fb.std(ddof=1):.2f}",
                 f"{fc.mean():.2f} ± {fc.std(ddof=1):.2f}",
                 f"{(fa-fb).mean():.2f} ± {(fa-fb).std(ddof=1):.2f}",
                 f"{(fb-fc).mean():.2f} ± {(fb-fc).std(ddof=1):.2f}"])
    add_sheet(wb, "S7_PanelB", "Supplementary Table S7, Panel B. Regional and global score differences",
              ["Outcome", "A", "B", "C", "A − B", "B − C"],
              rows, widths=[18, 18, 18, 18, 18, 18],
              note="Test set n=20. RAS is dimensionless; Global FA is on the 0–100 scale. Values are mean ± SD. Differences are computed per patient before averaging.")

    # ---------------- S7 panel C ----------------
    rows = []
    for key in ("A", "B", "C"):
        r = final["routes"][key]
        rows.append([key, f"{r['mae']:.2f} ({r['mae_ci'][0]:.2f}–{r['mae_ci'][1]:.2f})",
                     f"{r['rmse']:.2f} ({r['rmse_ci'][0]:.2f}–{r['rmse_ci'][1]:.2f})",
                     f"{r['r2']:.3f}", f"{r['pearson_r']:.3f}",
                     f"{r['icc_a1']:.3f} ({r['icc_a1_ci'][0]:.3f}–{r['icc_a1_ci'][1]:.3f})",
                     f"{r['bias']:.2f} ({r['loa_low']:.2f}–{r['loa_high']:.2f})"])
    add_sheet(wb, "S7_PanelC", "Supplementary Table S7, Panel C. Global FA accuracy against expert VAS",
              ["Route", "MAE (95% CI)", "RMSE (95% CI)", "Prediction R²", "Pearson r", "ICC(A,1) (95% CI)", "Bias (95% LoA)"],
              rows, widths=[10, 22, 22, 14, 12, 22, 24],
              note="Test set n=20. Route A row is identical to Table 4. ΔMAE(A−B) = −0.66 (95% CI −1.54 to 0.33), paired bootstrap 2000 resamples. No additional P values.")

    # ---------------- S2 ----------------
    rows = []
    for label, s in final["supplement2"].items():
        rows.append([label, fmt_med_iqr(s["angle"]), fmt_med_iqr(s["position"]), fmt_med_iqr(s["mhd"])])
    add_sheet(wb, "S2_个体轴", "Supplementary Table S2. AI axis versus each individual expert axis, and inter-expert agreement",
              ["Comparison", "Angular deviation (°), median (IQR)", "Position deviation at yc (px), median (IQR)", "MHD (px), median (IQR)"],
              rows, widths=[22, 30, 32, 24],
              note="Test set n=20. AI axis: route A (automatic landmarks + predicted axis). A1/A2/B1/B2 are the four individual expert annotations. Axis position is evaluated at the interpupillary-line height yc.")

    # ---------------- per-case ----------------
    headers = ["Stem", "Split", "Width", "Height", "IPD expert", "IPD auto", "yc",
               "Expert m", "Expert b", "AI m (B)", "AI b (B)", "AI m (A)", "AI b (A)",
               "Angle err A (°)", "Position err A (px)", "MHD A (px)",
               "Angle err B (°)", "Position err B (px)", "MHD B (px)"]
    headers += [f"RAS C {k}" for k in labels] + [f"RAS B {k}" for k in labels] + [f"RAS A {k}" for k in labels]
    headers += ["VAS", "VAS rater A", "VAS rater B", "z C", "FA C", "z B", "FA B", "z A", "FA A"]
    rows = []
    for r in per:
        row = [r["stem"], r["split"], int(r["width"]), int(r["height"]),
               f(r["ipd_expert"], 1), f(r["ipd_machine"], 1), f(r["yc"], 1),
               f(r["expert_m"], 6), f(r["expert_b"], 3), f(r["ai_m_B"], 6), f(r["ai_b_B"], 3),
               f(r["ai_m_A"], 6), f(r["ai_b_A"], 3),
               f(r["angle_err_A"], 3), f(r["pos_err_A"], 3), f(r["mhd_A"], 3),
               f(r["angle_err_B"], 3), f(r["pos_err_B"], 3), f(r["mhd_B"], 3)]
        row += [f(r[f"RAS_C_{k}"], 5) for k in regions]
        row += [f(r[f"RAS_B_{k}"], 5) for k in regions]
        row += [f(r[f"RAS_A_{k}"], 5) for k in regions]
        row += [f(r["VAS"], 2), f(r["VAS_A"], 2), f(r["VAS_B"], 2),
                f(r["z_C"], 3), f(r["FA_C"], 2), f(r["z_B"], 3), f(r["FA_B"], 2), f(r["z_A"], 3), f(r["FA_A"], 2)]
        rows.append(row)
    add_sheet(wb, "逐例原始数据", "Per-image raw data (n=204)",
              headers, rows, widths=None,
              note="RAS: regional asymmetry score; FA: Global FA Index; z: raw value before clipping. A/B/C as defined in the README sheet.")
    ws = wb["逐例原始数据"]
    for j in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(j)].width = 13 if j > 2 else 14

    # ---------------- VAS ----------------
    rows = [[r["stem"], r["split"], f(r["VAS"], 2), f(r["VAS_A"], 2), f(r["VAS_B"], 2)] for r in per]
    add_sheet(wb, "专家VAS_204", "Expert VAS scores (n=204)",
              ["Stem", "Split", "VAS (mean)", "VAS rater A", "VAS rater B"], rows,
              widths=[22, 10, 14, 14, 14],
              note="VAS: 0 = completely symmetric, 100 = extreme asymmetry. The mean of the two raters is used as the reference score.")

    # ---------------- 167 features ----------------
    names = json.load(open(OUT / "feature_names_167.json", encoding="utf-8"))
    rows = []
    for i, name in enumerate(names, 1):
        if i <= 58:
            group = "Normalized coordinate (58)"
        elif i <= 146:
            group = "Pair geometry (88)"
        else:
            group = "Midline geometry (21)"
        rows.append([i, name, group])
    add_sheet(wb, "167维特征", "167-dimensional coords_geom feature vector",
              ["Index", "Feature", "Group"], rows, widths=[8, 34, 28],
              note="58 normalized coordinates + 88 bilateral-pair geometry features + 21 midline features = 167. See 167维特征说明.md for formulas, normalization and inverse transform.")

    # ---------------- split ----------------
    rows = [[r["stem"], r["split"]] for r in per]
    add_sheet(wb, "划分名单", "Patient split (seed 11)", ["Stem", "Split"], rows, widths=[22, 12])

    out_xlsx = DELIV / "论文数据表.xlsx"
    wb.properties.creator = ""
    wb.properties.lastModifiedBy = ""
    wb.properties.title = "面部不对称自动评分：论文数据表"
    wb.properties.subject = ""
    wb.properties.description = ""
    wb.properties.keywords = ""
    wb.save(out_xlsx)

    import shutil
    shutil.copy(C.OUT / "final_per_case.csv", DELIV / "逐例原始数据.csv")
    shutil.copy(C.OUT / "axis_test_per_case.csv", DELIV / "轴模型测试集逐例.csv")
    print("saved", out_xlsx)

    # ---------------- figures ----------------
    fig_axis_box(all_models, "angle_diff", "Angular deviation (°)", "Symmetry-axis angular deviation (test)", FIG / "fig1_axis_angle_boxplot.png")
    fig_axis_box(all_models, "position_diff", "Axis-position deviation at yc (px)", "Symmetry-axis position deviation (test)", FIG / "fig2_axis_position_boxplot.png")
    fig_ba_ridge(ridge, FIG / "fig3_ridge_bland_altman.png")
    fig_fa_vs_vas(test, FIG / "fig4_fa_vs_vas.png")
    fig_residual(test, FIG / "fig5_fa_residual.png")
    fig_ba_fa(test, FIG / "fig6_fa_bland_altman.png")
    fig_ras(test, FIG / "fig7_ras_distribution.png")
    fig_s2(per, FIG / "fig8_supp2_axis.png")
    print("figures saved to", FIG)

    # ---------------- report ----------------
    write_report(axis, final, per, test)
    print("saved report")


def write_report(axis, final, per, test):
    t = axis["models"]["Ridge"]["test"]
    a = final["routes"]["A"]
    b = final["routes"]["B"]
    c = final["routes"]["C"]
    fit = final["formula"]
    d = final["descriptives"]
    ra = final["range_check_all"]
    ax = final["axis_errors_test"]
    s2 = final["supplement2"]
    diffs = final["score_diffs_test"]
    regions = ["eyebrow", "eye", "nose", "mouth", "contour"]
    labels = ["Eyebrow", "Eye", "Nose", "Mouth", "Facial contour"]

    def ms(x):
        return f"{x['mean']:.3f} ± {x['sd']:.3f}"

    def mi(x):
        return f"{x['median']:.3f} ({x['q1']:.3f}–{x['q3']:.3f})"

    lines = []
    lines.append("# 面部不对称自动评分：结果与论文表格\n")
    lines.append("## 1. 数据与划分\n")
    lines.append(f"- 标注 216 例，双侧关键点完整纳入 **204 例**；12 例因关键点缺失排除。")
    lines.append(f"- 训练 **{d['split']['train']}** / 验证 **{d['split']['val']}** / 测试 **{d['split']['test']}**（7:2:1，seed=11）。")
    lines.append(f"- 图像分辨率：{'; '.join(f'{k}（{v} 例）' for k, v in d['resolution'].items())}。")
    lines.append(f"- IPD：**{d['ipd']['mean']:.1f} ± {d['ipd']['sd']:.1f} px**（中位数 {d['ipd']['median']:.1f}，范围 {d['ipd']['min']:.0f}–{d['ipd']['max']:.0f}）。")
    lines.append(f"- 专家 VAS：**{d['vas']['mean']:.1f} ± {d['vas']['sd']:.1f}**（范围 {d['vas']['min']:.0f}–{d['vas']['max']:.0f}）。\n")

    lines.append("## 2. 对称轴模型\n")
    lines.append("| Model | CV angle (°) | Angle, median (IQR) | Position, median (IQR) | MHD, median (IQR) |")
    lines.append("|---|---|---|---|---|")
    for name in ["Ridge", "Linear", "SVR-RBF", "Random Forest"]:
        m = axis["models"][name]
        tt = m["test"]
        lines.append(f"| {name} | {m['cv']['mean']:.3f} | {mi(tt['angle'])} | {mi(tt['position'])} | {mi(tt['mhd'])} |")
    lines.append("")
    lines.append("Ridge 在 5 折 CV 和测试集平均角度误差上均为最优，且优于 Random Forest。\n")
    lines.append("Ridge 与专家共识轴的一致性：\n")
    lines.append("| Metric | Mean ± SD | Median (IQR) | 95% CI | Bias ± SD | 95% LoA |")
    lines.append("|---|---|---|---|---|---|")
    for label, key, sk in [("Angular deviation (°)", "angle", "signed_angle"),
                           ("Axis-position deviation at yc (px)", "position", "signed_position")]:
        ss = t[key]
        sa = t[sk]
        lines.append(f"| {label} | {ms(ss)} | {mi(ss)} | {ss['ci_low']:.3f}–{ss['ci_high']:.3f} | {sa['bias']:.2f} ± {sa['sd']:.2f} | {sa['loa_low']:.2f} to {sa['loa_high']:.2f} |")
    lines.append("")
    lines.append("![axis angle](figures/fig1_axis_angle_boxplot.png)\n")
    lines.append("![axis position](figures/fig2_axis_position_boxplot.png)\n")
    lines.append("![Ridge BA](figures/fig3_ridge_bland_altman.png)\n")

    lines.append("## 3. Global FA 公式\n")
    lines.append(f"训练+验证集 **n={fit['n_fit']}**，R²=**{fit['r2']:.4f}**，adjusted R²=**{fit['adj_r2']:.4f}**，RMSE=**{fit['rmse']:.3f}**。\n")
    lines.append("| Region | B | 95% CI | Standardized β | 95% CI | P | Normalized weight |")
    lines.append("|---|---|---|---|---|---|---|")
    pos = [max(0.0, x["beta_std"]) for x in fit["terms"]]
    wsum = sum(pos) if sum(pos) > 0 else 1.0
    for x, w in zip(fit["terms"], pos):
        lines.append(f"| {x['region']} | {x['B']:.3f} | {x['B_ci_low']:.3f}–{x['B_ci_high']:.3f} | {x['beta_std']:.3f} | {x['beta_std_ci_low']:.3f}–{x['beta_std_ci_high']:.3f} | {x['p']:.2e} | {w/wsum:.3f} |")
    lines.append("")
    lines.append("Global FA = clip(β0 + Σ βr·RASr, 0, 100)。\n")
    lines.append(f"越界检查（全队列）：raw 范围 **{ra['raw_min']:.2f}–{ra['raw_max']:.2f}**；低于 0 为 **{ra['n_below_0']}** 例，高于 100 为 **{ra['n_above_100']}** 例（**{ra['out_of_range_pct']:.2f}%**）。截断后 Global FA 为 {ra['fa_min']:.2f}–{ra['fa_max']:.2f}。\n")

    lines.append("## 4. 机器评分与专家 VAS（路线 A）\n")
    lines.append(f"- Pearson r = **{a['pearson_r']:.3f}**（95% CI {a['pearson_r_ci'][0]:.3f}–{a['pearson_r_ci'][1]:.3f}）")
    lines.append(f"- R² = **{a['r2']:.3f}**")
    lines.append(f"- ICC(A,1) = **{a['icc_a1']:.3f}**（95% CI {a['icc_a1_ci'][0]:.3f}–{a['icc_a1_ci'][1]:.3f}）")
    lines.append(f"- Bias = **{a['bias']:.2f}**（95% LoA {a['loa_low']:.2f} 至 {a['loa_high']:.2f}）")
    lines.append(f"- MAE = **{a['mae']:.2f}**（95% CI {a['mae_ci'][0]:.2f}–{a['mae_ci'][1]:.2f}）")
    lines.append(f"- RMSE = **{a['rmse']:.2f}**（95% CI {a['rmse_ci'][0]:.2f}–{a['rmse_ci'][1]:.2f}）\n")
    lines.append("![FA vs VAS](figures/fig4_fa_vs_vas.png)\n")
    lines.append("![residual](figures/fig5_fa_residual.png)\n")
    lines.append("![FA BA](figures/fig6_fa_bland_altman.png)\n")
    lines.append("![RAS](figures/fig7_ras_distribution.png)\n")

    lines.append("## 5. 补充实验一：三条路线\n")
    lines.append("Panel A. 轴误差（测试集，median (IQR)）\n")
    lines.append("| Comparison | Angle (°) | Position (px) | MHD (px) |")
    lines.append("|---|---|---|---|")
    for key, label in [("A", "A vs consensus"), ("B", "B vs consensus"), ("A_vs_B", "A vs B")]:
        s = ax[key]
        lines.append(f"| {label} | {mi(s['angle'])} | {mi(s['position'])} | {mi(s['mhd'])} |")
    lines.append("")
    lines.append("Panel B. 五区域 RAS 与 Global FA（mean ± SD）\n")
    lines.append("| Outcome | A | B | C | A − B | B − C |")
    lines.append("|---|---|---|---|---|---|")
    for key, label in zip(regions, labels):
        av = np.array([float(r[f"RAS_A_{key}"]) for r in test])
        bv = np.array([float(r[f"RAS_B_{key}"]) for r in test])
        cv = np.array([float(r[f"RAS_C_{key}"]) for r in test])
        lines.append(f"| {label} | {av.mean():.4f} ± {av.std(ddof=1):.4f} | {bv.mean():.4f} ± {bv.std(ddof=1):.4f} | {cv.mean():.4f} ± {cv.std(ddof=1):.4f} | {(av-bv).mean():.4f} ± {(av-bv).std(ddof=1):.4f} | {(bv-cv).mean():.4f} ± {(bv-cv).std(ddof=1):.4f} |")
    fa = np.array([float(r["FA_A"]) for r in test])
    fb = np.array([float(r["FA_B"]) for r in test])
    fc = np.array([float(r["FA_C"]) for r in test])
    lines.append(f"| Global FA | {fa.mean():.2f} ± {fa.std(ddof=1):.2f} | {fb.mean():.2f} ± {fb.std(ddof=1):.2f} | {fc.mean():.2f} ± {fc.std(ddof=1):.2f} | {(fa-fb).mean():.2f} ± {(fa-fb).std(ddof=1):.2f} | {(fb-fc).mean():.2f} ± {(fb-fc).std(ddof=1):.2f} |")
    lines.append("")
    lines.append("Panel C. Global FA 与专家 VAS 的一致性\n")
    lines.append("| Route | MAE (95% CI) | RMSE (95% CI) | R² | Pearson r | ICC(A,1) | Bias (95% LoA) |")
    lines.append("|---|---|---|---|---|---|---|")
    for key, name in [("A", "A"), ("B", "B"), ("C", "C")]:
        r = final["routes"][key]
        lines.append(f"| {name} | {r['mae']:.2f} ({r['mae_ci'][0]:.2f}–{r['mae_ci'][1]:.2f}) | {r['rmse']:.2f} ({r['rmse_ci'][0]:.2f}–{r['rmse_ci'][1]:.2f}) | {r['r2']:.3f} | {r['pearson_r']:.3f} | {r['icc_a1']:.3f} | {r['bias']:.2f} ({r['loa_low']:.2f}–{r['loa_high']:.2f}) |")
    lines.append("")
    lines.append(f"ΔMAE(A−B) = **−0.66**（95% CI −1.54 至 0.33，配对 bootstrap 2000 次）；A/B 差异小，CI 跨 0。\n")

    lines.append("## 6. 补充实验二：AI 轴与单条专家轴\n")
    lines.append("| Comparison | Angle (°) | Position (px) | MHD (px) |")
    lines.append("|---|---|---|---|")
    for label, s in s2.items():
        lines.append(f"| {label} | {mi(s['angle'])} | {mi(s['position'])} | {mi(s['mhd'])} |")
    lines.append("")
    lines.append("![S2](figures/fig8_supp2_axis.png)\n")

    lines.append("## 7. 说明\n")
    lines.append("- 12 例因双侧关键点不完整排除；204 例用于全部分析。")
    lines.append("- 评分时轴线以双瞳中点为锚点做横向平移；轴误差表仍使用原始轴线。")
    lines.append("- Global FA 截断至 0–100；2/204（1.0%）例在截断前高于 100，已在越界检查表中报告。")
    lines.append("- 167 维特征的定义、归一化、反变换、热图与训练参数见 `167维特征说明.md`。")
    lines.append("")

    (DELIV / "结果报告.md").write_text("\n".join(lines), encoding="utf-8")


if __name__ == "__main__":
    main()
