#!/usr/bin/env python3
"""Export comparison report to Excel workbook per 对比模型补充数据清单.docx structure."""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

import numpy as np

REPO_ROOT = Path(__file__).resolve().parents[2]
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

from src.landmark.schema import POINT_NAMES_31, REGIONS, NUM_POINTS

SDR_THRESHOLDS = [5, 10, 15, 20, 25, 30, 40]
REGION_DISPLAY = {"jaw": "Facial contour", "eyebrow": "Eyebrow", "eye": "Eye", "nose": "Nose", "mouth": "Mouth"}
REGION_ORDER = ["eyebrow", "eye", "nose", "mouth", "jaw"]

HEADER_FILL = PatternFill("solid", fgColor="4472C4")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TH_FILL = PatternFill("solid", fgColor="D9E1F2")
TH_FONT = Font(bold=True)
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row, c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER


def style_th(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row, c)
        cell.fill = TH_FILL
        cell.font = TH_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER


def set_col_widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w


def write_table2(ws, t2):
    ws.title = "表2 总体性能"
    headers = ["Model", "N", "MRE (px) Mean±SD", "MRE Median", "MRE IQR", "MRE 95%CI",
               "NME (%IPD) Mean±SD", "SDR@5%", "SDR@10%", "SDR@15%", "SDR@20%", "SDR@25%", "SDR@30%", "SDR@40%", "SDR AUC"]
    ws.append(headers)
    style_header(ws, 1, len(headers))
    for r in t2:
        ws.append([
            r["model"], r["n"],
            f"{r['mre_mean_px']:.2f} ± {r['mre_sd_px']:.2f}",
            f"{r['mre_median_px']:.2f}",
            f"{r['mre_iqr'][0]:.2f}–{r['mre_iqr'][1]:.2f}",
            f"{r['mre_ci95'][0]:.2f}–{r['mre_ci95'][1]:.2f}",
            f"{r['nme_mean_pct']:.2f} ± {r['nme_sd_pct']:.2f}",
            f"{r['sdr_px']['5']*100:.1f}", f"{r['sdr_px']['10']*100:.1f}", f"{r['sdr_px']['15']*100:.1f}",
            f"{r['sdr_px']['20']*100:.1f}", f"{r['sdr_px']['25']*100:.1f}", f"{r['sdr_px']['30']*100:.1f}",
            f"{r['sdr_px']['40']*100:.1f}", f"{r['sdr_auc']:.1f}",
        ])
    set_col_widths(ws, [16, 6, 20, 12, 16, 16, 22, 9, 9, 9, 9, 9, 9, 9, 9])


def write_table3(ws, t3, model_names):
    ws.title = "表3 区域性能"
    headers = ["Region"] + [f"{m} MRE(px)" for m in model_names] + \
              [f"{m} NME(%IPD)" for m in model_names] + \
              [f"{m} SDR@15%" for m in model_names] + \
              [f"{m} SDR@30%" for m in model_names]
    ws.append(headers)
    style_header(ws, 1, len(headers))
    for region in REGION_ORDER:
        row = [REGION_DISPLAY[region]]
        for m in model_names:
            rd = t3[m].get(region, {})
            row.append(f"{rd.get('mre_mean_px', 0):.2f}")
        for m in model_names:
            rd = t3[m].get(region, {})
            row.append(f"{rd.get('nme_mean_pct', 0):.2f}")
        for m in model_names:
            rd = t3[m].get(region, {})
            sdr = rd.get("sdr_px", {})
            row.append(f"{float(sdr.get('15', 0))*100:.1f}")
        for m in model_names:
            rd = t3[m].get(region, {})
            sdr = rd.get("sdr_px", {})
            row.append(f"{float(sdr.get('30', 0))*100:.1f}")
        ws.append(row)
    set_col_widths(ws, [16] + [16] * len(model_names) + [14] * len(model_names) + [14] * len(model_names) + [14] * len(model_names))


def write_table4(ws, t4, model_names):
    ws.title = "表4 逐点误差"
    headers = ["Region", "Landmark"]
    for m in model_names:
        headers += [f"{m} MRE(px)", f"{m} SDR@15%"]
    ws.append(headers)
    style_header(ws, 1, len(headers))
    for region in REGION_ORDER:
        for idx in REGIONS[region]:
            name = POINT_NAMES_31[idx]
            row = [REGION_DISPLAY[region], name]
            for m in model_names:
                pt = t4[m][name]
                row.append(f"{pt['mre_px']:.2f}" if pt["mre_px"] is not None else "-")
                row.append(f"{pt['sdr15']*100:.1f}" if pt["sdr15"] is not None else "-")
            ws.append(row)
    set_col_widths(ws, [16, 12] + [16, 14] * len(model_names))


def write_table5(ws, t5):
    ws.title = "表5 对称轴"
    headers = ["指标", "Linear Regression", "SVR-RBF", "Random Forest", "Ridge"]
    ws.append(headers)
    style_header(ws, 1, len(headers))
    rows = [
        ("Angular deviation (°), Mean±SD", "angular", "mean_sd"),
        ("Angular deviation, Median(IQR)", "angular", "median_iqr"),
        ("Angular deviation, 95% CI", "angular", "ci95"),
        ("Signed angular bias (°)", "angular", "signed_bias"),
        ("Angular 95% LoA (°)", "angular", "loa"),
        ("Axis-position deviation at yc (px), Mean±SD", "position", "mean_sd"),
        ("Position deviation, Median(IQR)", "position", "median_iqr"),
        ("Position deviation, 95% CI", "position", "ci95"),
        ("Signed position bias (px)", "position", "signed_bias"),
        ("Position 95% LoA (px)", "position", "loa"),
        ("Normalized position error (%IPD)", "position", "pct_ipd"),
        ("Mean horizontal distance (px)", "mhd", "mean_sd"),
        ("Mean horizontal distance (%IPD)", "mhd", "pct_ipd"),
    ]
    order = ["Linear", "SVR-RBF", "RandomForest", "Ridge"]
    for label, group, key in rows:
        vals = []
        for m in order:
            v = t5.get(m, {}).get(group, {}).get(key)
            if key == "signed_bias" and v is not None:
                vals.append(f"{v:.2f}")
            elif key == "pct_ipd" and v is not None:
                vals.append(f"{v:.2f}%")
            elif v is not None:
                vals.append(str(v))
            else:
                vals.append("-")
        ws.append([label] + vals)
    set_col_widths(ws, [34, 22, 22, 22, 22])


def write_table6(ws, t6, model_names):
    ws.title = "表6 复杂度"
    headers = ["指标"] + model_names
    ws.append(headers)
    style_header(ws, 1, len(headers))
    metric_rows = [
        ("Parameters（M）", "params_m"),
        ("GFLOPs", "gflops"),
        ("Model size（MB）", "model_size_mb"),
        ("Peak memory（MB）", "peak_memory_mb"),
        ("Preprocessing（ms）", "preprocess_ms"),
        ("Landmark inference（ms）", "inference_ms"),
        ("End-to-end runtime（ms）", "e2e_ms"),
        ("FPS", "fps"),
    ]
    for label, key in metric_rows:
        ws.append([label] + [t6.get(m, {}).get(key, "-") for m in model_names])
    set_col_widths(ws, [24] + [14] * len(model_names))


def write_table7(ws, t7):
    ws.title = "表7 训练条件"
    if not t7:
        return
    model_names = [k for k in t7[0].keys() if k != "metric"]
    headers = ["项目"] + model_names
    ws.append(headers)
    style_header(ws, 1, len(headers))
    for row in t7:
        ws.append([row["metric"]] + [row[m] for m in model_names])
    set_col_widths(ws, [24] + [22] * len(model_names))


def write_manifest(ws, lm_dirs, ax_dirs):
    ws.title = "数据量核对"
    ws.append(["任务", "样本量", "模型数", "说明"])
    style_header(ws, 1, 4)
    for name, mdir in lm_dirs.items():
        with open(Path(mdir) / "paper_metrics_summary.json") as f:
            s = json.load(f)
        ws.append(["关键点检测", f"{s['n_samples']}张×29点", 1, f"{name}"])
    for name, adir in ax_dirs.items():
        with open(Path(adir) / "test_sample_metrics.json") as f:
            sm = json.load(f)
        ws.append(["对称轴预测", f"{len(sm)}张", 1, f"{name}"])
    set_col_widths(ws, [18, 20, 10, 30])


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--report", required=True, help="path to report_data.json")
    parser.add_argument("--axis", nargs="+", required=True, help="name:path axis model dirs")
    parser.add_argument("--lm", nargs="+", required=True, help="name:path landmark model dirs")
    parser.add_argument("--output", required=True)
    args = parser.parse_args()

    with open(args.report) as f:
        data = json.load(f)
    t2 = data["table2_overall"]
    t3 = data["table3_region"]
    t4 = data["table4_per_point"]
    t5 = data["table5_axis"]
    t6 = data["table6_complexity"]
    t7 = data.get("table7_training", [])

    lm_dirs = {spec.split(":")[0]: Path(spec.split(":")[1]) for spec in args.lm}
    ax_dirs = {spec.split(":")[0]: Path(spec.split(":")[1]) for spec in args.axis}

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws2 = wb.create_sheet()
    write_table2(ws2, t2)
    write_table3(wb.create_sheet(), t3, [r["model"] for r in t2])
    write_table4(wb.create_sheet(), t4, [r["model"] for r in t2])
    write_table5(wb.create_sheet(), t5)
    write_table6(wb.create_sheet(), t6, [r["model"] for r in t2])
    write_table7(wb.create_sheet(), t7)
    write_manifest(wb.create_sheet(), lm_dirs, ax_dirs)

    out = Path(args.output)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print(f"Excel saved to {out}")
    print("Sheets:", wb.sheetnames)


if __name__ == "__main__":
    main()
